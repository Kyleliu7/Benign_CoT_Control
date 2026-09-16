import openpyxl
import json
from pathlib import Path
import numpy as np

PKG_DIR = Path(__file__).resolve().parent.parent
wb_path = PKG_DIR / "results" / "ALL_EXPERIMENTS_BY_MODEL_AUDITED.xlsx"
wb = openpyxl.load_workbook(wb_path)

bench_path = PKG_DIR / "data" / "eval_prompts" / "reasonif_dataset_300.json"
with open(bench_path, "r", encoding="utf-8") as f:
    bench = json.load(f)
c_map = {r["dataset_index"]: r["constraint_name"][0] for r in bench}

file_mapping = {
    7: "reasonif_qwen3_14b_base_untouched.jsonl",
    8: "reasonif_qwen3_14b_prefix_ack_requests.jsonl",
    9: "reasonif_qwen3_14b_prefix_constraint_off.jsonl",
    10: "reasonif_qwen3_14b_prefix_constraint_on.jsonl",
    11: "reasonif_qwen3_14b_sft_gpt52_high.jsonl",
    12: "reasonif_qwen3_14b_sft_gpt52_long.jsonl",
    13: "reasonif_qwen3_14b_sft_claude_37.jsonl",
    14: "reasonif_qwen3_14b_sft_qwen3_235b.jsonl",
    15: "reasonif_qwen3_14b_sft_reasonflux.jsonl",
    16: "reasonif_qwen3_14b_sft_gpt52_gen.jsonl",
    17: "reasonif_qwen3_14b_sft_output_mask.jsonl",
    18: "reasonif_qwen3_14b_sft_think_mask.jsonl",
    19: "reasonif_qwen3_14b_sft_no_reasoning.jsonl",
    20: "reasonif_qwen3_14b_sft_thinking_false.jsonl",
    21: "reasonif_qwen3_14b_sft_self_distill.jsonl",
    22: "reasonif_qwen3_14b_sft_svamp_meta.jsonl",
    23: "reasonif_gpt_oss_20b_base.jsonl",
    24: "reasonif_gpt_oss_20b_lora.jsonl",
    25: "reasonif_qwen3_14b_sft_gpt52_norm.jsonl",
}

cols = [
    "change_case:english_capital",
    "detectable_format:json_format",
    "language:reasoning_language",
    "length_constraint_checkers:number_words",
    "punctuation:no_comma",
    "startend:end_checker",
]

sheet = wb["ReasonIF - Master Summary"]
for r_idx, fname in file_mapping.items():
    fpath = PKG_DIR / "results" / "scored_runs" / fname
    with open(fpath, "r", encoding="utf-8") as f:
        rows = [json.loads(l) for l in f]
    tokens = [r.get("output_tokens", 0) or len(r.get("raw_output", "").split()) * 1.3 for r in rows]
    mean_tok = round(float(np.mean(tokens)), 2)
    sheet.cell(r_idx, 8, mean_tok)

    for _ci, c in enumerate(cols):
        c_rows = [r for r in rows if c_map.get(r.get("dataset_index")) == c]
        vals = [1 if (r.get("official_instruction_following") if r.get("official_instruction_following") is not None else r.get("instruction_following")) else 0 for r in c_rows]
        comp = float(np.mean(vals)) if vals else 0.0
        sheet.cell(r_idx, 9 + _ci, round(comp, 4))

if "Haskins - Qwen3 Pref-OFF" in wb.sheetnames:
    wb["Haskins - Qwen3 Pref-OFF"].cell(5, 2, "Dynamic Teacher Constraint-OFF Prefixes (51 unique task-specific prefixes, k=10 tokens; e.g. \"**Thinking Process:**\\n\\n\")")

if "Haskins - Qwen3 Pref-ON" in wb.sheetnames:
    wb["Haskins - Qwen3 Pref-ON"].cell(5, 2, "Dynamic Teacher Constraint-ON Prefixes (188 unique task-specific prefixes, k=10 tokens; teacher-steered constraint acknowledgment)")

wb.save(wb_path)
print("SUCCESS: Workbook updated!")
